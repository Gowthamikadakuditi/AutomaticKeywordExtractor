"""
Unit tests for Document Reader utility (PDF, DOCX, TXT).
"""

import io
import unittest
import docx
import pymupdf
from utils.document_reader import extract_text_from_stream, DocumentReaderError

class TestDocumentReader(unittest.TestCase):

    def test_txt_extraction(self):
        sample_txt = "Artificial Intelligence and Machine Learning project report.".encode("utf-8")
        text, filename = extract_text_from_stream("sample.txt", sample_txt)
        self.assertEqual(filename, "sample.txt")
        self.assertIn("Artificial Intelligence", text)

    def test_docx_extraction(self):
        # Create an in-memory docx
        doc = docx.Document()
        doc.add_paragraph("Natural language processing with Python and Flask.")
        doc.add_paragraph("Automatic keyword extraction test case.")
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        docx_bytes = bio.read()

        text, filename = extract_text_from_stream("paper.docx", docx_bytes)
        self.assertEqual(filename, "paper.docx")
        self.assertIn("Natural language processing", text)
        self.assertIn("Automatic keyword extraction", text)

    def test_pdf_extraction(self):
        # Create an in-memory PDF using PyMuPDF
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page()
        page.insert_text((50, 72), "Deep Learning and Keyphrase Extraction Research Paper")
        pdf_bytes = pdf_doc.tobytes()
        pdf_doc.close()

        text, filename = extract_text_from_stream("research.pdf", pdf_bytes)
        self.assertEqual(filename, "research.pdf")
        self.assertIn("Deep Learning and Keyphrase Extraction", text)

    def test_empty_pdf_detection(self):
        # Create a PDF page with no text (simulates blank page / scanned image)
        pdf_doc = pymupdf.open()
        pdf_doc.new_page()
        pdf_bytes = pdf_doc.tobytes()
        pdf_doc.close()

        with self.assertRaises(DocumentReaderError) as ctx:
            extract_text_from_stream("empty.pdf", pdf_bytes)
        self.assertIn("no extractable digital text", str(ctx.exception))

    def test_unsupported_file_extension(self):
        with self.assertRaises(DocumentReaderError) as ctx:
            extract_text_from_stream("malicious.exe", b"malicious content")
        self.assertIn("Unsupported file format", str(ctx.exception))

    def test_empty_file_bytes(self):
        with self.assertRaises(DocumentReaderError):
            extract_text_from_stream("", b"")

if __name__ == "__main__":
    unittest.main()
